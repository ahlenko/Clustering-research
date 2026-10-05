from docx import Document
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
p='research_report/Звіт_з_дослідницької_практики_доповнений.docx';d=Document(p)
for x in list(d.paragraphs):
 if not x.text and x._p.xpath('.//w:br[@w:type="page"]'):
  nxt=x._p.getnext()
  if nxt is not None and nxt.tag==qn('w:p'):
   from docx.text.paragraph import Paragraph
   Paragraph(nxt,x._parent).paragraph_format.page_break_before=True
   x._p.getparent().remove(x._p)
short=[
'Під час дослідницької практики обґрунтовано методику порівняння кластеризації текстів за внутрішньою структурою груп, відповідністю еталонним темам і витратами ресурсів. Сформульовано гіпотези щодо багатостартового пошуку, паралелізму, числа кластерів і мовної обробки; визначено умови їх перевірки.',
'У контрольному експерименті на 50 українських документах агломеративний метод отримав найбільші Silhouette (0,0106) і F1 macro (0,3029). Однак низькі значення Silhouette та близькі до нуля ARI свідчать про слабку відокремленість і тематичну узгодженість. DBSCAN із заданими параметрами відніс увесь корпус до шуму.',
'Parallel K-Means із TF-IDF поліпшив частину внутрішніх критеріїв, але не F1 та ARI; час був у 8,60 раза більшим за базовий K-Means. Порівняння різної кількості стартів не встановлює ефекту паралелізму. Рекомендація K = 7 за п’яти тем підтвердила розбіжність геометричної та тематичної структури.',
'Практичний внесок охоплює інтеграцію інструмента, організацію паралельних стартів і аналіз результатів. Порівняння з науковими працями обґрунтовує подальшу перевірку семантичних подань, але не дає підстав заявляти про перевищення результатів інших авторів на відмінних даних.',
'Мету практики досягнуто в частині методики та оцінювання наявних конфігурацій. Узагальнення потребує незалежних корпусів, перевірки лематизації, серій різних ініціалізацій і повних протоколів часу та пам’яті. Вибір методу має враховувати якість, охоплення документів і ресурси, зокрема результати, що не підтверджують гіпотези.'
]
inside=False;i=0
for x in d.paragraphs:
 if x.style.name=='Heading 1' and x.text=='Висновки':inside=True;continue
 if inside and i<5 and x.text:
  x.text=short[i];i+=1
 if inside and x.style.name=='Heading 1':break
# Keep the compact final comparison table together.
t=d.tables[-1]
for row in t.rows[:-1]:
 for c in row.cells:
  for x in c.paragraphs:x.paragraph_format.keep_with_next=True
d.save(p)
