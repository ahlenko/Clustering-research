from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
p=Path('practice_report/Звіт_з_практики_уточнений.docx');d=Document(p)
for q in d.paragraphs:
 if 'Початковий запуск із лематизацією' in q.text:
  q.text='Запуск із лематизацією завершився винятком ValueError через відсутність українського словника pymorphy3. Пакет pymorphy3-dicts-uk зазначено у requirements.txt, проте в перевіреному середовищі його не встановлено. Контрольний експеримент виконано без лематизації; після встановлення словника цей етап потрібно перевірити повторно.'
 if 'Графік середнього часу кластеризації' in q.text:
  q.text='Графік середнього часу кластеризації наведено на рис. 7. Показник clustering_seconds охоплює виконання алгоритму, але не підготовку корпусу, завантаження моделі, векторизацію, рекомендацію K чи побудову графіків. Тому він не дорівнює загальній тривалості експерименту.'
 if q.text:
  for r in q.runs:
   r.text=r.text.replace('візуального побудування','візуальної побудови').replace('за запуску з параметром','під час запуску з параметром')
for t in d.tables:
 for row in t.rows:
  for c in row.cells:
   for q in c.paragraphs:
    for r in q.runs:r.text=r.text.replace('Відсутня стовпець','Відсутній стовпець')
# Use page breaks attached to the following heading, avoiding blank pages.
paras=list(d.paragraphs)
for i,q in enumerate(paras):
 if q.text.strip():continue
 breaks=q._p.findall('.//'+qn('w:br'))
 if any(b.get(qn('w:type'))=='page' for b in breaks):
  for nxt in paras[i+1:]:
   if nxt.text.strip():
    nxt.paragraph_format.page_break_before=True;break
  q._p.getparent().remove(q._p)
d.save(p)
